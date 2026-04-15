from __future__ import annotations

from html import escape
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from backend.app.schemas.resume import (
    ResumeAnalyzeRequest,
    ResumeAutoDraftRequest,
    ResumeAutoFileRequest,
    ResumeDraftCreateRequest,
    ResumeDraftResponse,
    ResumeDraftUpdateRequest,
    ResumeExportRequest,
    ResumeFileGenerateRequest,
    ResumeGenerateRequest,
    ResumePreviewResponse,
    ResumeResponse,
    ResumeTemplateOption,
)
from backend.app.services.ai_chat_service import AiChatService
from backend.app.services.in_memory_store import database
from backend.app.services.media_service import MediaService
from backend.app.services.profile_service import ProfileService
from backend.app.schemas.media import PhotoStudioRequest
from backend.app.schemas.profile import UserProfileCreate


class ResumeService:
    EXPORT_DIR = Path(__file__).resolve().parents[3] / "tmp_exports"
    TEMPLATE_LIBRARY = [
        {
            "key": "modern_sidebar",
            "name": "Modern Sidebar",
            "description": "Tek sayfa, güçlü sidebar ve teknik profil ağırlıklı premium düzen.",
            "layout_hint": "Sidebar'da skills, diller, sertifikalar; ana alanda özet, deneyim ve projeler.",
        },
        {
            "key": "executive_grid",
            "name": "Executive Grid",
            "description": "Üst düzey ve kurumsal başvurular için daha dengeli ve ağırbaşlı görünüm.",
            "layout_hint": "Daha kurumsal bloklar, net başlık yapısı ve güçlü özet alanı.",
        },
        {
            "key": "minimal_ats",
            "name": "Minimal ATS",
            "description": "ATS uyumunu ve okunabilirliği öne çıkaran sade metin öncelikli düzen.",
            "layout_hint": "Az görsel yük, net bölüm hiyerarşisi ve düz güçlü içerik.",
        },
        {
            "key": "edinburgh_profile",
            "name": "Edinburgh Profile",
            "description": "Referans görseldeki gibi güçlü sol panel, temiz zaman çizgisi ve editorial CV görünümü.",
            "layout_hint": "Solda kişisel detaylar ve puanlı yetkinlikler, sağda kurumsal section akışı.",
        },
        {
            "key": "royal_timeline",
            "name": "Royal Timeline",
            "description": "Mavi panel ve zaman çizgili, daha klasik ama premium CV görünümü.",
            "layout_hint": "Fotoğraflı sol panel, sağda timeline odaklı deneyim akışı.",
        },
        {
            "key": "classic_timeline",
            "name": "Classic Timeline",
            "description": "Temiz beyaz zemin, sade kurumsal çizgi ve timeline odaklı yerleşim.",
            "layout_hint": "Az renk, net başlıklar ve okunaklı klasik CV düzeni.",
        },
        {
            "key": "burgundy_panel",
            "name": "Burgundy Panel",
            "description": "Koyu bordo yan panel ve güçlü başlık şeritleriyle dikkat çeken düzen.",
            "layout_hint": "Solda profil, sağda bloklu içerik ve sıcak premium tonlar.",
        },
        {
            "key": "crimson_edge",
            "name": "Crimson Edge",
            "description": "Sağda renkli vurgu paneliyle modern ve net kurumsal görünüm.",
            "layout_hint": "Temiz beyaz gövde, sağ tarafta kırmızı vurgu alanı.",
        },
        {
            "key": "onyx_accent",
            "name": "Onyx Accent",
            "description": "Siyah panel ve sıcak vurgu rengiyle iddialı modern CV tasarımı.",
            "layout_hint": "Solda koyu profil paneli, sağda vurgulu hero ve temiz içerik akışı.",
        },
    ]

    @classmethod
    def list_templates(cls) -> list[ResumeTemplateOption]:
        return [ResumeTemplateOption(**item) for item in cls.TEMPLATE_LIBRARY]

    @staticmethod
    def _resolve_template_key(template: str | None) -> str:
        known = {
            "modern_sidebar",
            "executive_grid",
            "minimal_ats",
            "edinburgh_profile",
            "royal_timeline",
            "classic_timeline",
            "burgundy_panel",
            "crimson_edge",
            "onyx_accent",
        }
        if template not in known:
            return "modern_sidebar"
        return template

    @staticmethod
    def _stringify_certifications(certifications: list | None) -> list[str]:
        if not certifications:
            return []
        normalized: list[str] = []
        for item in certifications:
            if isinstance(item, str):
                normalized.append(item)
                continue
            if isinstance(item, dict):
                line = " - ".join(
                    str(part) for part in [item.get("name"), item.get("issuer"), item.get("year")] if part
                )
                if line:
                    normalized.append(line)
        return normalized

    @classmethod
    def _extract_resume_blocks(cls, draft: dict) -> tuple[dict[str, str], str, str, str, str]:
        profile = draft["profile"]
        section_map = cls._split_sections(draft["content"], profile)
        summary_text = section_map.pop("PROFESSIONAL SUMMARY", None) or section_map.pop("PROFESYONEL ÖZET", None) or draft["summary"]
        experience_text = section_map.pop("WORK EXPERIENCE", None) or section_map.pop("EXPERIENCE", None) or cls._build_experience_text(profile)
        project_text = section_map.pop("KEY PROJECTS", None) or section_map.pop("PROJECTS", None) or cls._build_project_text(profile)
        education_text = section_map.pop("EDUCATION", None) or section_map.pop("EĞİTİM", None) or cls._build_education_text(profile)
        return section_map, summary_text, experience_text, project_text, education_text

    @staticmethod
    def _build_experience_text(profile: dict) -> str:
        experiences = profile.get("experiences") or []
        blocks: list[str] = []
        for item in experiences:
            date_range = " - ".join(part for part in [item.get("start_date"), item.get("end_date")] if part)
            header_parts = [item.get("title"), item.get("company")]
            if item.get("location"):
                header_parts.append(item.get("location"))
            header = " | ".join(part for part in header_parts if part)
            if date_range:
                header = f"{header} | {date_range}" if header else date_range
            if header:
                blocks.append(header)
            for bullet in item.get("bullets") or []:
                blocks.append(f"- {bullet}")
            tech = item.get("technologies") or []
            if tech:
                blocks.append(f"Technologies: {', '.join(tech)}")
            blocks.append("")
        return "\n".join(blocks).strip()

    @staticmethod
    def _build_project_text(profile: dict) -> str:
        projects = profile.get("projects") or []
        blocks: list[str] = []
        for item in projects:
            header = " | ".join(part for part in [item.get("name"), item.get("subtitle")] if part)
            if header:
                blocks.append(header)
            for bullet in item.get("bullets") or []:
                blocks.append(f"- {bullet}")
            tech = item.get("technologies") or []
            if tech:
                blocks.append(f"Technologies: {', '.join(tech)}")
            if item.get("link"):
                blocks.append(str(item["link"]))
            blocks.append("")
        return "\n".join(blocks).strip()

    @classmethod
    def _build_education_text(cls, profile: dict) -> str:
        education = profile.get("education") or []
        certifications = cls._stringify_certifications(profile.get("certifications"))
        blocks: list[str] = []
        for item in education:
            date_range = " - ".join(part for part in [item.get("start_date"), item.get("end_date")] if part)
            header = " | ".join(part for part in [item.get("degree"), item.get("school"), item.get("location")] if part)
            if date_range:
                header = f"{header} | {date_range}" if header else date_range
            if header:
                blocks.append(header)
            for detail in item.get("details") or []:
                blocks.append(f"- {detail}")
            blocks.append("")
        if certifications:
            blocks.extend(certifications)
        return "\n".join(blocks).strip()

    @staticmethod
    def analyze_resume(payload: ResumeAnalyzeRequest) -> ResumeResponse:
        result = AiChatService.analyze_resume(payload.content, payload.target_role)
        return ResumeResponse(
            title=result["title"],
            summary=result["summary"],
            content=result["content"],
            recommendations=result["recommendations"],
            template=None,
            template_reason=None,
        )

    @staticmethod
    def generate_resume(payload: ResumeGenerateRequest) -> ResumeResponse:
        result = AiChatService.generate_resume(
            payload.profile.model_dump(),
            payload.target_role,
            payload.template,
            payload.tone,
            payload.generation_context,
        )
        return ResumeResponse(
            title=result["title"],
            summary=result["summary"],
            content=result["content"],
            recommendations=result["recommendations"],
            template=result.get("template") or payload.template,
            template_reason=(payload.generation_context or {}).get("template_reason"),
        )

    @classmethod
    def create_preview(cls, payload: ResumeGenerateRequest) -> ResumePreviewResponse:
        result = cls.generate_resume(payload)
        preview_id = uuid4().hex
        html = cls._render_preview_html(
            profile=payload.profile.model_dump(),
            result=result,
            target_role=payload.target_role,
        )
        database["resume_previews"][preview_id] = {
            "html": html,
            "title": result.title,
        }
        return ResumePreviewResponse(
            preview_id=preview_id,
            preview_url=f"/api/v1/resumes/preview/{preview_id}",
            title=result.title,
        )

    @staticmethod
    def get_preview_html(preview_id: str) -> str | None:
        preview = database["resume_previews"].get(preview_id)
        if not preview:
            return None
        return str(preview["html"])

    @classmethod
    def create_draft(cls, payload: ResumeDraftCreateRequest) -> ResumeDraftResponse:
        result = cls.generate_resume(payload)
        draft_id = uuid4().hex
        preview_id = cls._store_preview(
            profile=payload.profile.model_dump(),
            result=result,
            target_role=payload.target_role,
        )
        database["resume_drafts"][draft_id] = {
            "draft_id": draft_id,
            "profile": payload.profile.model_dump(),
            "target_role": payload.target_role,
            "generation_context": payload.generation_context or {},
            "title": result.title,
            "summary": result.summary,
            "content": result.content,
            "recommendations": result.recommendations,
            "profile_warnings": payload.profile_warnings,
            "template": result.template or payload.template,
            "template_reason": (payload.generation_context or {}).get("template_reason"),
            "tone": payload.tone,
            "finalized": False,
            "preview_id": preview_id,
        }
        return cls._draft_to_response(database["resume_drafts"][draft_id])

    @classmethod
    def create_auto_draft(cls, payload: ResumeAutoDraftRequest) -> ResumeDraftResponse | None:
        profile = ProfileService.get_profile(payload.profile_id)
        if not profile:
            return None

        profile_data = profile.model_dump()
        generation_context = dict(payload.generation_context or {})
        warnings = AiChatService.review_profile_for_resume(
            profile_data,
            payload.target_role or profile_data.get("target_role"),
            generation_context,
        )

        if payload.auto_professionalize_photo and profile_data.get("photo_url"):
            photo_style = generation_context.get("photo_style") or generation_context.get("visual_preferences") or "professional dark background"
            try:
                processed = MediaService.professionalize_photo(
                    PhotoStudioRequest(
                        photo_path=str(profile_data["photo_url"]),
                        template=payload.template,
                        style=photo_style,
                    )
                )
                profile_data["photo_url"] = processed.processed_path
                warnings.append("Yüklenen fotoğraf otomatik olarak CV’ye uygun profesyonel headshot görünümüne işlendi.")
            except FileNotFoundError:
                warnings.append("Profil fotoğrafı bulundu ama işlenebilir dosya çözümlenemedi; mevcut foto kullanılmadan devam edildi.")
        else:
            warnings.append("Profil fotoğrafı yoksa sistem yine CV üretir; premium template’lerde foto opsiyonel kalır.")

        generation_context["resume_review_notes"] = " | ".join(warnings[:4])

        return cls.create_draft(
            ResumeDraftCreateRequest(
                profile=UserProfileCreate.model_validate(profile_data),
                template=payload.template,
                tone=payload.tone,
                target_role=payload.target_role or profile_data.get("target_role"),
                generation_context=generation_context,
                profile_warnings=warnings,
            )
        )

    @classmethod
    def get_draft(cls, draft_id: str) -> ResumeDraftResponse | None:
        draft = database["resume_drafts"].get(draft_id)
        if not draft:
            return None
        return cls._draft_to_response(draft)

    @classmethod
    def update_draft(
        cls,
        draft_id: str,
        payload: ResumeDraftUpdateRequest,
    ) -> ResumeDraftResponse | None:
        draft = database["resume_drafts"].get(draft_id)
        if not draft:
            return None

        updates = payload.model_dump(exclude_none=True)
        for key, value in updates.items():
            draft[key] = value

        draft["profile_warnings"] = AiChatService.review_draft_for_resume(
            draft["profile"],
            draft["summary"],
            draft["content"],
            draft.get("target_role"),
            draft.get("generation_context") or {},
        )

        draft["preview_id"] = cls._store_preview(
            profile=draft["profile"],
            result=ResumeResponse(
                title=draft["title"],
                summary=draft["summary"],
                content=draft["content"],
                recommendations=draft["recommendations"],
                template=draft["template"],
                template_reason=draft.get("template_reason"),
            ),
            target_role=draft.get("target_role"),
        )
        return cls._draft_to_response(draft)

    @classmethod
    def finalize_draft(cls, draft_id: str) -> ResumeDraftResponse | None:
        draft = database["resume_drafts"].get(draft_id)
        if not draft:
            return None
        draft["finalized"] = True
        return cls._draft_to_response(draft)

    @classmethod
    def export_draft(cls, draft_id: str, payload: ResumeExportRequest) -> Path | None:
        draft = database["resume_drafts"].get(draft_id)
        if not draft:
            return None

        cls.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        safe_name = cls._slugify(draft["profile"].get("full_name") or "cv")
        output_path = cls.EXPORT_DIR / f"{safe_name}-{draft_id}.{payload.format}"

        if payload.format == "docx":
            cls._build_docx(draft, output_path)
        else:
            cls._build_pdf(draft, output_path)

        database["resume_exports"][draft_id] = {
            "path": str(output_path),
            "format": payload.format,
        }
        return output_path

    @classmethod
    def generate_file(cls, payload: ResumeFileGenerateRequest) -> Path:
        draft = cls.create_draft(
            ResumeDraftCreateRequest(
                profile=payload.profile,
                template=payload.template,
                tone=payload.tone,
                target_role=payload.target_role,
            )
        )
        cls.finalize_draft(draft.draft_id)
        export_path = cls.export_draft(
            draft.draft_id,
            ResumeExportRequest(format=payload.format),
        )
        if export_path is None:
            raise RuntimeError("Resume export could not be created")
        return export_path

    @classmethod
    def generate_auto_file(cls, payload: ResumeAutoFileRequest) -> Path | None:
        draft = cls.create_auto_draft(
            ResumeAutoDraftRequest(
                profile_id=payload.profile_id,
                template=payload.template,
                tone=payload.tone,
                target_role=payload.target_role,
                generation_context=payload.generation_context,
                auto_professionalize_photo=payload.auto_professionalize_photo,
            )
        )
        if draft is None:
            return None

        cls.finalize_draft(draft.draft_id)
        export_path = cls.export_draft(
            draft.draft_id,
            ResumeExportRequest(format=payload.format),
        )
        if export_path is None:
            raise RuntimeError("Resume export could not be created")
        return export_path

    @classmethod
    def get_draft_preview_html(cls, draft_id: str) -> str | None:
        draft = database["resume_drafts"].get(draft_id)
        if not draft:
            return None
        preview_id = draft.get("preview_id")
        if not preview_id:
            return None
        preview = database["resume_previews"].get(preview_id)
        if not preview:
            return None
        return str(preview["html"])

    @classmethod
    def _store_preview(
        cls,
        profile: dict,
        result: ResumeResponse,
        target_role: str | None,
    ) -> str:
        preview_id = uuid4().hex
        html = cls._render_preview_html(
            profile=profile,
            result=result,
            target_role=target_role,
        )
        database["resume_previews"][preview_id] = {
            "html": html,
            "title": result.title,
        }
        return preview_id

    @staticmethod
    def _draft_to_response(draft: dict) -> ResumeDraftResponse:
        return ResumeDraftResponse(
            draft_id=draft["draft_id"],
            title=draft["title"],
            summary=draft["summary"],
            content=draft["content"],
            recommendations=draft["recommendations"],
            profile_warnings=draft.get("profile_warnings", []),
            template=draft.get("template"),
            template_reason=draft.get("template_reason"),
            finalized=bool(draft.get("finalized")),
            preview_url=f"/api/v1/resumes/drafts/{draft['draft_id']}/preview",
        )


    @staticmethod
    def _render_preview_html(
        profile: dict,
        result: ResumeResponse,
        target_role: str | None,
    ) -> str:
        template_key = ResumeService._resolve_template_key(result.template)
        skills = profile.get("skills") or []
        languages = profile.get("languages") or []
        certifications = ResumeService._stringify_certifications(profile.get("certifications"))
        links = [
            profile.get("linkedin_url"),
            profile.get("github_url"),
            profile.get("portfolio_url"),
        ]
        safe_links = [escape(link) for link in links if link]
        photo_path = ResumeService._resolve_photo_path(profile.get("photo_url"))
        photo_html = (
            f"<img src='file:///{photo_path.replace(chr(92), '/')}' alt='Profile photo' class='avatar' />"
            if photo_path
            else ""
        )
        recommendation_html = "".join(
            f"<li>{escape(item)}</li>" for item in result.recommendations
        )
        skills_html = "".join(f"<span class='chip'>{escape(item)}</span>" for item in skills)
        languages_html = "".join(f"<li>{escape(item)}</li>" for item in languages)
        certifications_html = "".join(f"<li>{escape(item)}</li>" for item in certifications)
        links_html = "".join(
            f"<a href='{link}' target='_blank' rel='noreferrer'>{link}</a>" for link in safe_links
        )
        section_map = ResumeService._split_sections(result.content, profile)
        summary_text = escape(section_map.pop("PROFESSIONAL SUMMARY", None) or profile.get("summary") or result.summary)
        main_blocks = "".join(
            f"<div class='section'><h2>{escape(title)}</h2><pre>{escape(content)}</pre></div>"
            for title, content in section_map.items()
        )

        theme_map = {
            "modern_sidebar": {
                "bg": "linear-gradient(180deg, #eef3fb 0%, #f8fafc 100%)",
                "hero": "linear-gradient(135deg, #132441 0%, #1f4b99 55%, #315bb1 100%)",
                "paper": "#ffffff",
                "side_bg": "#f8faff",
                "line": "#d8deea",
                "accent": "#1f4b99",
                "layout": "split",
            },
            "executive_grid": {
                "bg": "linear-gradient(180deg, #f4f1eb 0%, #fbfaf7 100%)",
                "hero": "linear-gradient(135deg, #2a241e 0%, #4a3b2b 100%)",
                "paper": "#fffdf9",
                "side_bg": "#f5f0e8",
                "line": "#e1d7c9",
                "accent": "#6b4d2f",
                "layout": "stacked",
            },
            "minimal_ats": {
                "bg": "#f5f7fa",
                "hero": "linear-gradient(135deg, #20252d 0%, #2b313b 100%)",
                "paper": "#ffffff",
                "side_bg": "#ffffff",
                "line": "#e4e7ec",
                "accent": "#20252d",
                "layout": "single",
            },
            "edinburgh_profile": {
                "bg": "linear-gradient(180deg, #eeedf3 0%, #f7f7fa 100%)",
                "hero": "linear-gradient(135deg, #54557a 0%, #7072a1 100%)",
                "paper": "#ffffff",
                "side_bg": "#f4f4f7",
                "line": "#d7d8e1",
                "accent": "#54557a",
                "layout": "split",
            },
            "royal_timeline": {
                "bg": "linear-gradient(180deg, #edf2f8 0%, #f8fbff 100%)",
                "hero": "linear-gradient(135deg, #466b99 0%, #6f8eb8 100%)",
                "paper": "#ffffff",
                "side_bg": "#edf3fb",
                "line": "#d7dde7",
                "accent": "#466b99",
                "layout": "split",
            },
            "classic_timeline": {
                "bg": "#fafafa",
                "hero": "linear-gradient(135deg, #2b2b2b 0%, #555555 100%)",
                "paper": "#ffffff",
                "side_bg": "#ffffff",
                "line": "#d1d1d1",
                "accent": "#2b2b2b",
                "layout": "single",
            },
            "burgundy_panel": {
                "bg": "linear-gradient(180deg, #f7f0f3 0%, #fcfafb 100%)",
                "hero": "linear-gradient(135deg, #6b2c3e 0%, #8c4458 100%)",
                "paper": "#ffffff",
                "side_bg": "#f8f2f5",
                "line": "#e0cfd6",
                "accent": "#6b2c3e",
                "layout": "split",
            },
            "crimson_edge": {
                "bg": "#faf7f7",
                "hero": "linear-gradient(135deg, #ba2323 0%, #d14a4a 100%)",
                "paper": "#ffffff",
                "side_bg": "#fff5f5",
                "line": "#ead2d2",
                "accent": "#ba2323",
                "layout": "single",
            },
            "onyx_accent": {
                "bg": "linear-gradient(180deg, #f8f6f2 0%, #fdfcf9 100%)",
                "hero": "linear-gradient(135deg, #050505 0%, #2a2a2a 100%)",
                "paper": "#ffffff",
                "side_bg": "#f5f0e7",
                "line": "#e3dacb",
                "accent": "#a87124",
                "layout": "split",
            },
        }
        theme = theme_map[template_key]

        return f"""<!doctype html>
<html lang="tr">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>{escape(result.title)}</title>
  <style>
    :root {{
      color-scheme: light;
      --bg: {theme["bg"]};
      --paper: {theme["paper"]};
      --ink: #172033;
      --muted: #5c6881;
      --line: {theme["line"]};
      --accent: {theme["accent"]};
      --accent-soft: #edf3ff;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: "Segoe UI", Arial, sans-serif;
      background: var(--bg);
      color: var(--ink);
      padding: 32px 16px;
    }}
    .page {{
      max-width: 980px;
      margin: 0 auto;
      background: var(--paper);
      border: 1px solid var(--line);
      border-radius: 28px;
      overflow: hidden;
      box-shadow: 0 25px 80px rgba(23, 32, 51, 0.12);
    }}
    .hero {{
      padding: 40px;
      background: {theme["hero"]};
      color: white;
    }}
    .hero h1 {{
      margin: 0;
      font-size: 40px;
      line-height: 1.05;
    }}
    .hero p {{
      margin: 10px 0 0;
      color: rgba(255,255,255,0.88);
    }}
    .hero-top {{
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 18px;
    }}
    .avatar {{
      width: 92px;
      height: 92px;
      border-radius: 999px;
      object-fit: cover;
      border: 3px solid rgba(255,255,255,0.65);
      box-shadow: 0 10px 28px rgba(0,0,0,0.15);
      flex-shrink: 0;
    }}
    .meta {{
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      margin-top: 18px;
    }}
    .meta span, .chip {{
      display: inline-flex;
      align-items: center;
      border-radius: 999px;
      padding: 8px 12px;
      background: rgba(255,255,255,0.12);
      font-size: 13px;
    }}
    .content {{
      display: {"grid" if theme["layout"] == "split" else "block"};
      grid-template-columns: 1.4fr 0.8fr;
      gap: 0;
    }}
    .main, .side {{ padding: 32px; }}
    .side {{
      background: {theme["side_bg"]};
      border-left: {"1px solid var(--line)" if theme["layout"] == "split" else "0"};
    }}
    h2 {{
      margin: 0 0 14px;
      font-size: 15px;
      letter-spacing: 0.18em;
      text-transform: uppercase;
      color: var(--accent);
    }}
    .section + .section {{ margin-top: 28px; }}
    .summary {{
      font-size: 16px;
      line-height: 1.8;
      color: var(--ink);
    }}
    .card {{
      border: 1px solid var(--line);
      background: var(--accent-soft);
      border-radius: 18px;
      padding: 16px;
      line-height: 1.7;
    }}
    .chips {{
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
    }}
    .chips .chip {{
      background: white;
      border: 1px solid var(--line);
      color: var(--ink);
    }}
    ul {{
      margin: 0;
      padding-left: 18px;
      line-height: 1.8;
      color: var(--muted);
    }}
    pre {{
      margin: 0;
      white-space: pre-wrap;
      font-family: inherit;
      line-height: 1.8;
      color: var(--muted);
    }}
    .links a {{
      display: block;
      color: var(--accent);
      text-decoration: none;
      margin-bottom: 10px;
      word-break: break-all;
    }}
    @media (max-width: 800px) {{
      .content {{ grid-template-columns: 1fr; }}
      .side {{
        border-left: 0;
        border-top: {"1px solid var(--line)" if theme["layout"] == "split" else "0"};
      }}
      .hero, .main, .side {{
        padding: 24px;
      }}
      .hero h1 {{
        font-size: 32px;
      }}
    }}
  </style>
</head>
<body>
  <div class="page">
    <div class="hero">
      <div class="hero-top">
        <div>
          <h1>{escape(profile.get("full_name", "Aday"))}</h1>
          <p>{escape(target_role or profile.get("target_role") or "Profesyonel CV")}</p>
        </div>
        {photo_html}
      </div>
      <div class="meta">
        <span>{escape(profile.get("email", "-"))}</span>
        {f"<span>{escape(profile.get('phone'))}</span>" if profile.get("phone") else ""}
        {f"<span>{escape(profile.get('location'))}</span>" if profile.get("location") else ""}
      </div>
    </div>
    <div class="content">
      <div class="main">
        <div class="section">
          <h2>Profesyonel Özet</h2>
          <div class="summary">{summary_text}</div>
        </div>
        {main_blocks}
      </div>
      <aside class="side" style="display:{'block' if theme['layout'] != 'single' else 'none'}">
        <div class="section">
          <h2>Yetkinlikler</h2>
          <div class="chips">{skills_html or "<span class='chip'>Henüz eklenmedi</span>"}</div>
        </div>
        {f'<div class="section"><h2>Öneriler</h2><ul>{recommendation_html}</ul></div>' if result.recommendations else ''}
        <div class="section">
          <h2>Diller</h2>
          <ul>{languages_html or "<li>Belirtilmedi</li>"}</ul>
        </div>
        <div class="section">
          <h2>Sertifikalar</h2>
          <ul>{certifications_html or "<li>Belirtilmedi</li>"}</ul>
        </div>
        <div class="section links">
          <h2>Bağlantılar</h2>
          {links_html or "<span class='chip'>Bağlantı eklenmedi</span>"}
        </div>
      </aside>
    </div>
  </div>
</body>
</html>"""

    @classmethod
    def _build_docx(cls, draft: dict, output_path: Path) -> None:
        document = Document()
        profile = draft["profile"]
        template_key = cls._resolve_template_key(draft.get("template"))
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)
        skills = profile.get("skills") or []
        languages = profile.get("languages") or []
        certifications = cls._stringify_certifications(profile.get("certifications"))
        links = [value for value in (profile.get("linkedin_url"), profile.get("github_url"), profile.get("portfolio_url")) if value]
        photo_path = cls._resolve_photo_path(profile.get("photo_url"))

        section_colors = {
            "modern_sidebar": "1F4B99",
            "executive_grid": "6B4D2F",
            "minimal_ats": "111827",
            "edinburgh_profile": "5B5D89",
            "royal_timeline": "355C9A",
            "classic_timeline": "2F313F",
            "burgundy_panel": "7A3045",
            "crimson_edge": "B53A44",
            "onyx_accent": "C8924A",
        }
        accent = section_colors.get(template_key, "1F4B99")

        normal_style = document.styles["Normal"]
        normal_style.font.name = "Calibri"
        normal_style.font.size = Pt(10.5)
        title_style = document.styles["Title"]
        title_style.font.name = "Calibri"
        title_style.font.size = Pt(22)

        title = document.add_paragraph()
        title.style = "Title"
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run(profile.get("full_name") or "CV")
        title_run.bold = True

        if photo_path:
            photo_para = document.add_paragraph()
            photo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            photo_run = photo_para.add_run()
            photo_run.add_picture(photo_path, width=Inches(1.2))

        if profile.get("target_role"):
            role_para = document.add_paragraph()
            role_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            role_run = role_para.add_run(profile["target_role"])
            role_run.italic = True

        contact_parts = [profile.get("email"), profile.get("phone"), profile.get("location")]
        contact_line = " | ".join(part for part in contact_parts if part)
        if contact_line:
            contact_para = document.add_paragraph()
            contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact_para.add_run(contact_line)

        if links:
            links_para = document.add_paragraph()
            links_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            links_para.add_run(" | ".join(links))

        if draft.get("template_reason"):
            hint = document.add_paragraph()
            hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hint.add_run(f"Template kararı: {draft['template_reason']}").italic = True

        if draft.get("profile_warnings"):
            cls._add_docx_section(document, "Kontrol Edilmesi Gerekenler", draft["profile_warnings"], accent, bullets=True)

        cls._add_docx_section(document, "Profesyonel Özet", [summary_text], accent)
        if experience_text:
            cls._add_docx_section(document, "Deneyim", experience_text.splitlines(), accent)
        if project_text:
            cls._add_docx_section(document, "Projeler", project_text.splitlines(), accent)
        if education_text:
            cls._add_docx_section(document, "Eğitim ve Sertifikalar", education_text.splitlines(), accent)
        if skills:
            cls._add_docx_section(document, "Yetkinlikler", [", ".join(skills)], accent)
        if languages:
            cls._add_docx_section(document, "Diller", [", ".join(languages)], accent)
        if certifications:
            cls._add_docx_section(document, "Sertifikalar", certifications, accent, bullets=True)
        if draft["recommendations"]:
            cls._add_docx_section(document, "AI Önerileri", draft["recommendations"], accent, bullets=True)

        document.save(output_path)

    @staticmethod
    def _add_docx_section(
        document: Document,
        title: str,
        lines: list[str],
        accent: str,
        bullets: bool = False,
    ) -> None:
        cleaned = [line.strip() for line in lines if line and line.strip()]
        if not cleaned:
            return
        heading = document.add_paragraph()
        heading_run = heading.add_run(title)
        heading_run.bold = True
        heading_run.font.size = Pt(13)
        heading_run.font.color.rgb = RGBColor.from_string(accent)
        for line in cleaned:
            if bullets or line.startswith("- "):
                item = line[2:] if line.startswith("- ") else line
                document.add_paragraph(item, style="List Bullet")
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(line)

    @classmethod
    def _build_pdf(cls, draft: dict, output_path: Path) -> None:
        template_key = cls._resolve_template_key(draft.get("template"))
        font_name = cls._ensure_pdf_font()
        if template_key == "modern_sidebar":
            cls._build_pdf_modern_sidebar(draft, output_path, font_name)
            return
        if template_key == "burgundy_panel":
            cls._build_pdf_burgundy_panel(draft, output_path, font_name)
            return
        if template_key == "onyx_accent":
            cls._build_pdf_onyx_accent(draft, output_path, font_name)
            return
        if template_key == "royal_timeline":
            cls._build_pdf_royal_timeline(draft, output_path, font_name)
            return
        if template_key == "classic_timeline":
            cls._build_pdf_classic_timeline(draft, output_path, font_name)
            return
        if template_key == "crimson_edge":
            cls._build_pdf_crimson_edge(draft, output_path, font_name)
            return
        if template_key == "executive_grid":
            cls._build_pdf_executive_grid(draft, output_path, font_name)
            return
        if template_key == "edinburgh_profile":
            cls._build_pdf_edinburgh_profile(draft, output_path, font_name)
            return

        document = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
        )
        palette = {
            "modern_sidebar": {"title": "#172033", "accent": "#1f4b99", "muted": "#5c6881", "body": "#233048"},
            "executive_grid": {"title": "#2f261d", "accent": "#6b4d2f", "muted": "#7a6756", "body": "#382f27"},
            "minimal_ats": {"title": "#111827", "accent": "#111827", "muted": "#4b5563", "body": "#1f2937"},
        }[template_key]
        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                name="CvTitle",
                fontName=font_name,
                fontSize=23 if template_key == "executive_grid" else 21,
                leading=27 if template_key == "executive_grid" else 25,
                textColor=HexColor(palette["title"]),
                spaceAfter=10 if template_key != "minimal_ats" else 6,
            )
        )
        styles.add(
            ParagraphStyle(
                name="CvHeading",
                fontName=font_name,
                fontSize=12,
                leading=15,
                textColor=HexColor(palette["accent"]),
                spaceBefore=10 if template_key == "executive_grid" else 8,
                spaceAfter=6,
            )
        )
        styles.add(
            ParagraphStyle(
                name="CvBody",
                fontName=font_name,
                fontSize=10.3 if template_key == "minimal_ats" else 10.5,
                leading=15 if template_key != "minimal_ats" else 14,
                textColor=HexColor(palette["body"]),
                spaceAfter=5 if template_key == "minimal_ats" else 6,
            )
        )
        styles.add(
            ParagraphStyle(
                name="CvMeta",
                fontName=font_name,
                fontSize=10,
                leading=14,
                textColor=HexColor(palette["muted"]),
                spaceAfter=4,
            )
        )
        styles.add(
            ParagraphStyle(
                name="SidebarHeading",
                fontName=font_name,
                fontSize=10.5,
                leading=13,
                textColor=HexColor(palette["accent"]),
                spaceBefore=4,
                spaceAfter=4,
            )
        )
        styles.add(
            ParagraphStyle(
                name="SidebarBody",
                fontName=font_name,
                fontSize=9.5,
                leading=13,
                textColor=HexColor(palette["body"]),
                spaceAfter=3,
            )
        )

        profile = draft["profile"]
        story = [Paragraph(cls._escape_pdf(profile.get("full_name") or "CV"), styles["CvTitle"])]

        target_role = profile.get("target_role")
        if target_role:
            story.append(Paragraph(cls._escape_pdf(f"Hedef Rol: {target_role}"), styles["CvMeta"]))

        contact_parts = [profile.get("email"), profile.get("phone"), profile.get("location")]
        contact_line = " | ".join(part for part in contact_parts if part)
        if contact_line:
            story.append(Paragraph(cls._escape_pdf(contact_line), styles["CvMeta"]))

        story.append(Spacer(1, 8))

        section_map, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)
        sidebar_parts: list = []
        main_parts: list = []

        for heading, values in (
            ("TECHNICAL SKILLS", cls._group_skills(profile.get("skills") or [])),
            ("LANGUAGES", (profile.get("languages") or [])),
            ("CERTIFICATIONS", cls._stringify_certifications(profile.get("certifications"))),
            ("LINKS", cls._collect_links(profile)),
        ):
            if not values:
                continue
            sidebar_parts.append(Paragraph(heading, styles["SidebarHeading"]))
            if isinstance(values, dict):
                for label, items in values.items():
                    sidebar_parts.append(
                        Paragraph(cls._escape_pdf(f"{label}: {', '.join(items)}"), styles["SidebarBody"])
                    )
            else:
                for item in values:
                    sidebar_parts.append(Paragraph(cls._escape_pdf(str(item)), styles["SidebarBody"]))
            sidebar_parts.append(Spacer(1, 6))

        main_parts.append(Paragraph("PROFESSIONAL SUMMARY", styles["CvHeading"]))
        main_parts.append(Paragraph(cls._escape_pdf(summary_text), styles["CvBody"]))

        ordered_blocks = [
            ("WORK EXPERIENCE", experience_text),
            ("KEY PROJECTS", project_text),
            ("EDUCATION", education_text),
        ]
        for section_name, content in ordered_blocks:
            if not content:
                continue
            main_parts.append(Paragraph(section_name, styles["CvHeading"]))
            for paragraph in cls._normalize_content(content):
                main_parts.append(Paragraph(cls._escape_pdf(paragraph), styles["CvBody"]))

        for section_name, content in section_map.items():
            if section_name == "TECHNICAL SKILLS":
                continue
            main_parts.append(Paragraph(section_name, styles["CvHeading"]))
            for paragraph in cls._normalize_content(content):
                main_parts.append(Paragraph(cls._escape_pdf(paragraph), styles["CvBody"]))

        if sidebar_parts and template_key == "modern_sidebar":
            story.append(Paragraph("PROFILE SNAPSHOT", styles["CvHeading"]))
            story.extend(sidebar_parts)
            story.append(Spacer(1, 6))
        elif sidebar_parts and template_key == "executive_grid":
            story.append(Paragraph("CORE PROFILE", styles["CvHeading"]))
            story.extend(sidebar_parts)
            story.append(Spacer(1, 4))
        elif sidebar_parts and template_key == "minimal_ats":
            ats_skill_parts = [Paragraph("CORE SKILLS", styles["CvHeading"])]
            for group, values in cls._group_skills(profile.get("skills") or []).items():
                ats_skill_parts.append(
                    Paragraph(cls._escape_pdf(f"{group}: {', '.join(values)}"), styles["CvBody"])
                )
            main_parts = ats_skill_parts + [Spacer(1, 4)] + main_parts

        story.extend(main_parts)

        document.build(story)

    @classmethod
    def _build_pdf_modern_sidebar(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        sidebar_width = 62 * mm
        margin = 14 * mm
        main_left = sidebar_width + 14 * mm
        main_right = page_width - margin
        main_width = main_right - main_left

        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)

        role = profile.get("target_role") or draft.get("target_role") or ""

        def draw_shell(first_page: bool) -> tuple[float, float]:
            pdf.setFillColor(HexColor("#10233f"))
            pdf.rect(0, 0, sidebar_width, page_height, fill=1, stroke=0)
            sidebar_y = page_height - margin
            sidebar_y = cls._draw_sidebar_title(pdf, font_name, profile.get("full_name") or "Candidate", sidebar_y, sidebar_width)
            if role:
                sidebar_y = cls._draw_sidebar_text(pdf, font_name, role, sidebar_y - 2 * mm, sidebar_width, size=11, color="#dbe8ff")
            if first_page:
                photo_path = cls._resolve_photo_path(profile.get("photo_url"))
                sidebar_y = cls._draw_profile_photo(pdf, photo_path, sidebar_width / 2, sidebar_y - 12 * mm, 15 * mm, "#dbe8ff")
                sidebar_y -= 22 * mm
            return sidebar_y, page_height - margin

        y, main_y = draw_shell(first_page=True)

        contact_lines = [
            profile.get("email"),
            profile.get("phone"),
            profile.get("location"),
            profile.get("linkedin_url"),
            profile.get("github_url"),
            profile.get("portfolio_url"),
        ]
        contact_lines = [line for line in contact_lines if line]
        if contact_lines:
            y -= 8 * mm
            y = cls._draw_sidebar_section(pdf, font_name, "CONTACT", contact_lines, y, sidebar_width)

        skill_groups = cls._group_skills(profile.get("skills") or [])
        if skill_groups:
            y -= 2 * mm
            skill_lines = [f"{group}: {', '.join(items)}" for group, items in skill_groups.items()]
            y = cls._draw_sidebar_section(pdf, font_name, "TECH SKILLS", skill_lines, y, sidebar_width)

        languages = profile.get("languages") or []
        if languages:
            y -= 2 * mm
            y = cls._draw_sidebar_section(pdf, font_name, "LANGUAGES", languages, y, sidebar_width)

        certifications = cls._stringify_certifications(profile.get("certifications"))
        if certifications:
            y -= 2 * mm
            y = cls._draw_sidebar_section(pdf, font_name, "CERTIFICATIONS", certifications, y, sidebar_width)

        main_y = cls._draw_main_title(pdf, font_name, "PROFILE", main_y, main_left)
        main_y = cls._draw_main_paragraph(
            pdf,
            font_name,
            summary_text,
            main_left,
            main_y,
            main_width,
            reset_page=lambda: draw_shell(first_page=False)[1],
        )

        if experience_text:
            main_y -= 4 * mm
            main_y = cls._ensure_canvas_space(pdf, main_y, 24 * mm, lambda: draw_shell(first_page=False)[1])
            main_y = cls._draw_main_title(pdf, font_name, "WORK EXPERIENCE", main_y, main_left)
            main_y = cls._draw_main_block(
                pdf, font_name, experience_text, main_left, main_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
            )

        if project_text:
            main_y -= 4 * mm
            main_y = cls._ensure_canvas_space(pdf, main_y, 24 * mm, lambda: draw_shell(first_page=False)[1])
            main_y = cls._draw_main_title(pdf, font_name, "KEY PROJECTS", main_y, main_left)
            main_y = cls._draw_main_block(
                pdf, font_name, project_text, main_left, main_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
            )

        if education_text:
            main_y -= 4 * mm
            main_y = cls._ensure_canvas_space(pdf, main_y, 24 * mm, lambda: draw_shell(first_page=False)[1])
            main_y = cls._draw_main_title(pdf, font_name, "EDUCATION", main_y, main_left)
            cls._draw_main_block(
                pdf, font_name, education_text, main_left, main_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
            )

        pdf.save()

    @classmethod
    def _build_pdf_burgundy_panel(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        sidebar_width = 68 * mm
        margin = 14 * mm
        main_left = sidebar_width + 12 * mm
        main_width = page_width - main_left - margin
        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)
        role = profile.get("target_role") or draft.get("target_role") or ""

        def draw_shell(first_page: bool) -> tuple[float, float]:
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#6b2c3e"))
            pdf.rect(0, 0, sidebar_width, page_height, fill=1, stroke=0)
            if first_page:
                cls._draw_profile_photo(
                    pdf,
                    cls._resolve_photo_path(profile.get("photo_url")),
                    sidebar_width / 2,
                    page_height - 24 * mm,
                    13 * mm,
                    "#9f6674",
                )
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.setFont(font_name, 19)
            pdf.drawString(main_left, page_height - 18 * mm, profile.get("full_name") or "Candidate")
            pdf.setFillColor(HexColor("#4b202c"))
            pdf.setFont(font_name, 11)
            if role:
                pdf.drawString(main_left, page_height - 25 * mm, role)
            return page_height - 50 * mm, page_height - 38 * mm

        left_y, right_y = draw_shell(first_page=True)
        personal_lines = [line for line in [
            profile.get("location"),
            profile.get("phone"),
            profile.get("email"),
            profile.get("linkedin_url"),
            profile.get("github_url"),
            profile.get("portfolio_url"),
        ] if line]
        if personal_lines:
            left_y = cls._draw_burgundy_side_section(pdf, font_name, "KİŞİSEL BİLGİLER", personal_lines, left_y, sidebar_width)
        skills = (profile.get("skills") or [])[:5]
        if skills:
            left_y -= 4 * mm
            left_y = cls._draw_burgundy_bars(pdf, font_name, "BECERİLER", skills, left_y, sidebar_width)

        right_y = cls._draw_burgundy_card(pdf, font_name, "KİŞİSEL ÖZET", summary_text, main_left, right_y, main_width)
        right_y = cls._draw_burgundy_card(pdf, font_name, "İŞ DENEYİMİ", experience_text, main_left, right_y - 4 * mm, main_width, reset_page=lambda: draw_shell(False)[1])
        if education_text:
            right_y = cls._draw_burgundy_card(pdf, font_name, "EĞİTİM", education_text, main_left, right_y - 4 * mm, main_width, reset_page=lambda: draw_shell(False)[1])
        if project_text:
            cls._draw_burgundy_card(pdf, font_name, "PROJELER", project_text, main_left, right_y - 4 * mm, main_width, reset_page=lambda: draw_shell(False)[1])
        pdf.save()

    @classmethod
    def _build_pdf_onyx_accent(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        sidebar_width = 60 * mm
        margin = 14 * mm
        main_left = sidebar_width + 12 * mm
        main_width = page_width - main_left - margin
        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)
        role = profile.get("target_role") or draft.get("target_role") or ""

        def draw_shell(first_page: bool) -> tuple[float, float]:
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#050505"))
            pdf.rect(0, 0, sidebar_width, page_height, fill=1, stroke=0)
            if first_page:
                cls._draw_profile_photo(
                    pdf,
                    cls._resolve_photo_path(profile.get("photo_url")),
                    sidebar_width / 2,
                    page_height - 24 * mm,
                    13 * mm,
                    "#d7b27a",
                )
            pdf.setFillColor(HexColor("#ffc476"))
            pdf.roundRect(main_left, page_height - 28 * mm, main_width - 10 * mm, 16 * mm, 8 * mm, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#161616"))
            pdf.setFont(font_name, 21)
            pdf.drawString(main_left + 8 * mm, page_height - 18 * mm, profile.get("full_name") or "Candidate")
            pdf.setFont(font_name, 11)
            if role:
                pdf.drawString(main_left + 8 * mm, page_height - 24 * mm, role)
            return page_height - 56 * mm, page_height - 40 * mm

        left_y, right_y = draw_shell(first_page=True)
        personal_lines = [line for line in [profile.get("email"), profile.get("phone"), profile.get("location")] if line]
        if personal_lines:
            left_y = cls._draw_onyx_side_section(pdf, font_name, "KİŞİSEL BİLGİLER", personal_lines, left_y, sidebar_width)
        education_text_left = cls._build_education_text(profile)
        if education_text_left:
            left_y -= 4 * mm
            left_y = cls._draw_onyx_side_section(pdf, font_name, "EĞİTİM", cls._normalize_content(education_text_left)[:4], left_y, sidebar_width)
        skills = (profile.get("skills") or [])[:5]
        if skills:
            left_y -= 4 * mm
            left_y = cls._draw_burgundy_bars(pdf, font_name, "BECERİLER", skills, left_y, sidebar_width, bar_color="#ffffff", text_color="#ffffff", title_color="#ffffff")

        right_y = cls._draw_onyx_section(pdf, font_name, "KİŞİSEL ÖZET", summary_text, main_left, right_y, main_width)
        right_y = cls._draw_onyx_section(pdf, font_name, "İŞ DENEYİMİ", experience_text, main_left, right_y - 4 * mm, main_width, reset_page=lambda: draw_shell(False)[1])
        if project_text:
            right_y = cls._draw_onyx_section(pdf, font_name, "PROJELER", project_text, main_left, right_y - 4 * mm, main_width, reset_page=lambda: draw_shell(False)[1])
        if education_text:
            cls._draw_onyx_section(pdf, font_name, "REFERANSLAR / EĞİTİM", education_text, main_left, right_y - 4 * mm, main_width, reset_page=lambda: draw_shell(False)[1])
        pdf.save()

    @classmethod
    def _build_pdf_classic_timeline(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        margin = 15 * mm
        content_width = page_width - 2 * margin
        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)

        def draw_shell() -> float:
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#252525"))
            pdf.setFont(font_name, 22)
            pdf.drawString(margin, page_height - margin, profile.get("full_name") or "Candidate")
            pdf.setFillColor(HexColor("#4a4a4a"))
            pdf.setFont(font_name, 11)
            role = profile.get("target_role") or draft.get("target_role") or ""
            if role:
                pdf.drawString(margin, page_height - margin - 7 * mm, role)
            return page_height - margin - 18 * mm

        y = draw_shell()
        y = cls._draw_classic_header_meta(pdf, font_name, profile, margin, y, content_width)
        y = cls._draw_classic_section(pdf, font_name, "Profile", summary_text, margin, y, content_width)
        y = cls._draw_classic_timeline_section(pdf, font_name, "Work Experience", experience_text, margin, y, content_width, reset_page=draw_shell)
        if education_text:
            y = cls._draw_classic_timeline_section(pdf, font_name, "Education and Qualifications", education_text, margin, y, content_width, reset_page=draw_shell)
        if project_text:
            cls._draw_classic_section(pdf, font_name, "Projects", project_text, margin, y, content_width, reset_page=draw_shell)
        pdf.save()

    @classmethod
    def _build_pdf_crimson_edge(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        panel_width = 58 * mm
        margin = 14 * mm
        content_width = page_width - panel_width - (2 * margin)
        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)

        def draw_shell(first_page: bool) -> float:
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#ba2323"))
            pdf.rect(page_width - panel_width, 0, panel_width, page_height, fill=1, stroke=0)
            if first_page:
                cls._draw_profile_photo(
                    pdf,
                    cls._resolve_photo_path(profile.get("photo_url")),
                    margin + 11 * mm,
                    page_height - 18 * mm,
                    8 * mm,
                    "#f1d5d5",
                )
            pdf.setFillColor(HexColor("#1f1f1f"))
            pdf.setFont(font_name, 20)
            pdf.drawString(margin + 22 * mm, page_height - 14 * mm, profile.get("full_name") or "Candidate")
            pdf.setFont(font_name, 10.5)
            role = profile.get("target_role") or draft.get("target_role") or ""
            if role:
                pdf.drawString(margin + 22 * mm, page_height - 20 * mm, role)
            return page_height - 30 * mm

        y = draw_shell(first_page=True)
        side_x = page_width - panel_width + 8 * mm
        side_y = page_height - 32 * mm
        side_lines = [line for line in [profile.get("phone"), profile.get("email"), profile.get("location")] if line]
        if side_lines:
            side_y = cls._draw_crimson_side_section(pdf, font_name, "Kişisel Bilgiler", side_lines, side_x, side_y, panel_width - 16 * mm)
        skills = (profile.get("skills") or [])[:5]
        if skills:
            side_y -= 6 * mm
            cls._draw_crimson_skill_dots(pdf, font_name, "Beceriler", skills, side_x, side_y, panel_width - 16 * mm)

        y = cls._draw_crimson_section(pdf, font_name, "Kısa Özet", summary_text, margin, y, content_width, reset_page=lambda: draw_shell(False))
        y = cls._draw_crimson_section(pdf, font_name, "İş Deneyimi", experience_text, margin, y - 4 * mm, content_width, reset_page=lambda: draw_shell(False))
        if education_text:
            y = cls._draw_crimson_section(pdf, font_name, "Eğitim", education_text, margin, y - 4 * mm, content_width, reset_page=lambda: draw_shell(False))
        if project_text:
            cls._draw_crimson_section(pdf, font_name, "Referanslar / Projeler", project_text, margin, y - 4 * mm, content_width, reset_page=lambda: draw_shell(False))
        pdf.save()

    @classmethod
    def _build_pdf_executive_grid(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        margin = 15 * mm
        content_left = margin
        content_right = page_width - margin
        content_width = content_right - content_left
        left_col_width = 56 * mm
        gutter = 8 * mm
        right_col_x = content_left + left_col_width + gutter
        right_col_width = content_width - left_col_width - gutter

        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)

        role = profile.get("target_role") or draft.get("target_role") or "Executive Profile"

        def draw_shell() -> tuple[float, float, float]:
            pdf.setFillColor(HexColor("#f7f2ea"))
            pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.roundRect(margin, margin, content_width, page_height - (2 * margin), 6 * mm, fill=1, stroke=0)
            header_height = 34 * mm
            pdf.setFillColor(HexColor("#2f261d"))
            pdf.roundRect(margin, page_height - margin - header_height, content_width, header_height, 6 * mm, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.setFont(font_name, 24)
            pdf.drawString(content_left + 7 * mm, page_height - margin - 11 * mm, (profile.get("full_name") or "Candidate").upper())
            pdf.setFont(font_name, 10)
            pdf.setFillColor(HexColor("#eadcc8"))
            pdf.drawString(content_left + 7 * mm, page_height - margin - 18 * mm, role)
            meta_band_y = page_height - margin - header_height - 10 * mm
            pdf.setFillColor(HexColor("#6b4d2f"))
            pdf.roundRect(content_left + 7 * mm, meta_band_y, content_width - 14 * mm, 11 * mm, 3 * mm, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#fffaf4"))
            pdf.setFont(font_name, 9)
            meta_items = [profile.get("email"), profile.get("phone"), profile.get("location")]
            meta_text = "  •  ".join(item for item in meta_items if item) or "Premium career document"
            for line in cls._wrap_text(meta_text, font_name, 9, content_width - 22 * mm):
                pdf.drawString(content_left + 11 * mm, meta_band_y + 4.1 * mm, line)
                break
            return meta_band_y, content_left + left_col_width + (gutter / 2), meta_band_y - 8 * mm

        meta_band_y, divider_x, current_y = draw_shell()
        current_y = cls._draw_executive_title(pdf, font_name, "EXECUTIVE SUMMARY", current_y, content_left + 2 * mm)
        current_y = cls._draw_executive_paragraph(
            pdf,
            font_name,
            summary_text,
            content_left + 2 * mm,
            current_y,
            content_width - 4 * mm,
        )

        left_y = current_y - 7 * mm
        right_y = current_y - 7 * mm

        pdf.setStrokeColor(HexColor("#e2d8cc"))
        pdf.setLineWidth(0.8)
        pdf.line(divider_x, margin + 10 * mm, divider_x, left_y + 3 * mm)

        skill_groups = cls._group_skills(profile.get("skills") or [])
        if skill_groups:
            left_y = cls._draw_executive_title(pdf, font_name, "CORE CAPABILITIES", left_y, content_left)
            capability_lines = [f"{group}: {', '.join(items)}" for group, items in skill_groups.items()]
            left_y = cls._draw_executive_list(pdf, font_name, capability_lines, content_left, left_y, left_col_width, bullet=False)
            left_y -= 4 * mm

        certifications = cls._stringify_certifications(profile.get("certifications"))
        if certifications:
            left_y = cls._draw_executive_title(pdf, font_name, "CERTIFICATIONS", left_y, content_left)
            left_y = cls._draw_executive_list(pdf, font_name, certifications, content_left, left_y, left_col_width, bullet=False)
            left_y -= 4 * mm

        languages = profile.get("languages") or []
        if languages:
            left_y = cls._draw_executive_title(pdf, font_name, "LANGUAGES", left_y, content_left)
            left_y = cls._draw_executive_list(pdf, font_name, languages, content_left, left_y, left_col_width, bullet=False)
            left_y -= 4 * mm

        links = cls._collect_links(profile)
        if links:
            left_y = cls._draw_executive_title(pdf, font_name, "LINKS", left_y, content_left)
            left_y = cls._draw_executive_list(pdf, font_name, links, content_left, left_y, left_col_width, bullet=False)

        if experience_text:
            right_y = cls._ensure_canvas_space(pdf, right_y, 26 * mm, lambda: draw_shell()[2])
            right_y = cls._draw_executive_title(pdf, font_name, "PROFESSIONAL EXPERIENCE", right_y, right_col_x)
            right_y = cls._draw_executive_block(
                pdf, font_name, experience_text, right_col_x, right_y, right_col_width, reset_page=lambda: draw_shell()[2]
            )
            right_y -= 4 * mm

        if project_text:
            right_y = cls._ensure_canvas_space(pdf, right_y, 26 * mm, lambda: draw_shell()[2])
            right_y = cls._draw_executive_title(pdf, font_name, "SELECTED PROJECTS", right_y, right_col_x)
            right_y = cls._draw_executive_block(
                pdf, font_name, project_text, right_col_x, right_y, right_col_width, reset_page=lambda: draw_shell()[2]
            )
            right_y -= 4 * mm

        if education_text:
            right_y = cls._ensure_canvas_space(pdf, right_y, 26 * mm, lambda: draw_shell()[2])
            right_y = cls._draw_executive_title(pdf, font_name, "EDUCATION", right_y, right_col_x)
            cls._draw_executive_block(
                pdf, font_name, education_text, right_col_x, right_y, right_col_width, reset_page=lambda: draw_shell()[2]
            )

        pdf.save()

    @classmethod
    def _build_pdf_edinburgh_profile(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        sidebar_width = 68 * mm
        margin = 14 * mm
        main_left = sidebar_width + 14 * mm
        main_width = page_width - main_left - margin

        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)
        certifications = []

        def draw_shell(first_page: bool) -> tuple[float, float]:
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#f6f5f7"))
            pdf.rect(0, 0, sidebar_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#54557a"))
            pdf.rect(0, page_height - 54 * mm, sidebar_width, 54 * mm, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#6d6ea0"))
            pdf.circle(sidebar_width - 5 * mm, page_height - 42 * mm, 10 * mm, fill=1, stroke=0)
            name = (profile.get("full_name") or "Candidate").upper()
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.setFont(font_name, 16)
            for idx, line in enumerate(cls._wrap_text(name, font_name, 16, sidebar_width - 22 * mm)):
                pdf.drawCentredString(sidebar_width / 2, page_height - 12 * mm - (idx * 6.2 * mm), line)
            if first_page:
                cls._draw_profile_photo(
                    pdf,
                    cls._resolve_photo_path(profile.get("photo_url")),
                    sidebar_width / 2,
                    page_height - 33 * mm,
                    13 * mm,
                    "#d9dbe6",
                )
            return page_height - 68 * mm, page_height - 20 * mm

        left_y, right_y = draw_shell(first_page=True)
        personal_lines = [
            f"Role: {profile.get('target_role')}" if profile.get("target_role") else None,
            f"Email: {profile.get('email')}" if profile.get("email") else None,
            f"Phone: {profile.get('phone')}" if profile.get("phone") else None,
            f"Location: {profile.get('location')}" if profile.get("location") else None,
            profile.get("linkedin_url"),
            profile.get("github_url"),
            profile.get("portfolio_url"),
        ]
        personal_lines = [line for line in personal_lines if line]
        if personal_lines:
            left_y = cls._draw_edinburgh_side_section(pdf, font_name, "PERSONAL DETAILS", personal_lines, left_y, sidebar_width)

        left_y -= 2 * mm
        interests = ["Product Thinking", "System Design", "Performance", "Leadership"]
        left_y = cls._draw_edinburgh_side_section(pdf, font_name, "INTERESTS", interests, left_y, sidebar_width)

        left_y -= 2 * mm
        languages = profile.get("languages") or []
        if languages:
            left_y = cls._draw_edinburgh_rating_section(pdf, font_name, "LANGUAGES", languages, left_y, sidebar_width)

        left_y -= 2 * mm
        skill_items = (profile.get("skills") or [])[:5]
        if skill_items:
            left_y = cls._draw_edinburgh_rating_section(pdf, font_name, "SKILLS", skill_items, left_y, sidebar_width)

        right_y = cls._draw_edinburgh_intro(pdf, font_name, summary_text, main_left, right_y, main_width)
        right_y -= 2 * mm
        right_y = cls._draw_edinburgh_section(
            pdf, font_name, "WORK EXPERIENCE", experience_text, main_left, right_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
        )
        if project_text:
            right_y = cls._draw_edinburgh_section(
                pdf, font_name, "PROJECTS", project_text, main_left, right_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
            )

        edu_and_cert = education_text
        if certifications:
            extra = "\n".join(certifications)
            edu_and_cert = f"{education_text}\n{extra}".strip() if education_text else extra
        if edu_and_cert:
            right_y = cls._draw_edinburgh_section(
                pdf,
                font_name,
                "EDUCATION & QUALIFICATIONS",
                edu_and_cert,
                main_left,
                right_y,
                main_width,
                reset_page=lambda: draw_shell(first_page=False)[1],
            )

        pdf.save()

    @classmethod
    def _build_pdf_royal_timeline(cls, draft: dict, output_path: Path, font_name: str) -> None:
        pdf = canvas.Canvas(str(output_path), pagesize=A4)
        page_width, page_height = A4
        sidebar_width = 60 * mm
        margin = 14 * mm
        main_left = sidebar_width + 10 * mm
        main_width = page_width - main_left - margin

        profile = draft["profile"]
        _, summary_text, experience_text, project_text, education_text = cls._extract_resume_blocks(draft)
        certifications = cls._stringify_certifications(profile.get("certifications"))

        def draw_shell(first_page: bool) -> tuple[float, float]:
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.rect(0, 0, page_width, page_height, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#466b99"))
            pdf.rect(0, 0, sidebar_width, page_height, fill=1, stroke=0)
            if first_page:
                cls._draw_profile_photo(
                    pdf,
                    cls._resolve_photo_path(profile.get("photo_url")),
                    sidebar_width / 2,
                    page_height - 26 * mm,
                    17 * mm,
                    "#8fa9cb",
                )
            pdf.setFillColor(HexColor("#ffffff"))
            pdf.setFont(font_name, 18)
            name_y = page_height - 56 * mm
            for idx, line in enumerate(cls._wrap_text((profile.get("full_name") or "Candidate").upper(), font_name, 18, sidebar_width - 18 * mm)):
                pdf.drawString(8 * mm, name_y - idx * 7 * mm, line)
            return page_height - 84 * mm, page_height - 18 * mm

        left_y, right_y = draw_shell(first_page=True)
        personal_lines = [line for line in [
            f"Role: {profile.get('target_role')}" if profile.get("target_role") else None,
            profile.get("email"),
            profile.get("phone"),
            profile.get("location"),
            profile.get("linkedin_url"),
            profile.get("github_url"),
            profile.get("portfolio_url"),
        ] if line]
        if personal_lines:
            left_y = cls._draw_edinburgh_side_section(pdf, font_name, "PERSONAL DETAILS", personal_lines, left_y, sidebar_width)
        left_y -= 2 * mm
        languages = profile.get("languages") or []
        if languages:
            left_y = cls._draw_edinburgh_rating_section(pdf, font_name, "LANGUAGES", languages, left_y, sidebar_width)
        left_y -= 2 * mm
        skill_items = (profile.get("skills") or [])[:4]
        if skill_items:
            left_y = cls._draw_edinburgh_rating_section(pdf, font_name, "SKILLS", skill_items, left_y, sidebar_width)

        right_y = cls._draw_royal_intro(pdf, font_name, summary_text, main_left, right_y, main_width)
        right_y = cls._draw_royal_section(
            pdf, font_name, "WORK EXPERIENCE", experience_text, main_left, right_y - 1 * mm, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
        )
        if education_text:
            right_y = cls._draw_royal_section(
                pdf, font_name, "EDUCATION & QUALIFICATIONS", education_text, main_left, right_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
            )
        if project_text:
            right_y = cls._draw_royal_section(
                pdf, font_name, "PROJECTS", project_text, main_left, right_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
            )
        if certifications:
            right_y = cls._draw_royal_section(
                pdf, font_name, "CERTIFICATIONS", "\n".join(certifications), main_left, right_y, main_width, reset_page=lambda: draw_shell(first_page=False)[1]
            )
        pdf.save()

    @staticmethod
    def _normalize_content(content: str) -> list[str]:
        normalized = []
        for raw_line in content.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            normalized.append(line)
        return normalized or ["İçerik bulunamadı."]

    @staticmethod
    def _wrap_text(text: str, font_name: str, font_size: int, max_width: float) -> list[str]:
        words = text.split()
        if not words:
            return []
        lines: list[str] = []
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if pdfmetrics.stringWidth(candidate, font_name, font_size) <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
        return lines

    @classmethod
    def _draw_sidebar_title(cls, pdf: canvas.Canvas, font_name: str, text: str, y: float, sidebar_width: float) -> float:
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont(font_name, 22)
        for line in cls._wrap_text(text.upper(), font_name, 22, sidebar_width - 18 * mm):
            pdf.drawString(8 * mm, y, line)
            y -= 9 * mm
        return y

    @classmethod
    def _draw_sidebar_text(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        text: str,
        y: float,
        sidebar_width: float,
        size: int = 10,
        color: str = "#ffffff",
    ) -> float:
        pdf.setFillColor(HexColor(color))
        pdf.setFont(font_name, size)
        for line in cls._wrap_text(text, font_name, size, sidebar_width - 18 * mm):
            pdf.drawString(8 * mm, y, line)
            y -= 5.2 * mm
        return y

    @classmethod
    def _draw_sidebar_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        lines: list[str],
        y: float,
        sidebar_width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#c5d6f7"))
        pdf.setFont(font_name, 10)
        pdf.drawString(8 * mm, y, title)
        y -= 6 * mm
        for line in lines:
            y = cls._draw_sidebar_text(pdf, font_name, line, y, sidebar_width, size=9, color="#ffffff")
            y -= 1.5 * mm
        return y

    @staticmethod
    def _draw_main_title(pdf: canvas.Canvas, font_name: str, title: str, y: float, x: float) -> float:
        pdf.setFillColor(HexColor("#1f4b99"))
        pdf.setFont(font_name, 13)
        pdf.drawString(x, y, title)
        return y - 6 * mm

    @classmethod
    def _draw_main_paragraph(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        pdf.setFillColor(HexColor("#223047"))
        pdf.setFont(font_name, 10)
        for line in cls._wrap_text(text, font_name, 10, width):
            if reset_page is not None:
                y = cls._ensure_canvas_space(pdf, y, 24 * mm, reset_page)
            pdf.drawString(x, y, line)
            y -= 5.2 * mm
        return y

    @classmethod
    def _draw_main_block(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        pdf.setFillColor(HexColor("#223047"))
        current_size = 10
        pdf.setFont(font_name, current_size)
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if line.startswith("- "):
                bullet_lines = cls._wrap_text(line[2:], font_name, 10, width - 6 * mm)
                for idx, part in enumerate(bullet_lines):
                    if reset_page is not None:
                        y = cls._ensure_canvas_space(pdf, y, 24 * mm, reset_page)
                    prefix = "• " if idx == 0 else "  "
                    pdf.drawString(x, y, prefix + part)
                    y -= 5 * mm
                continue
            wrapped = cls._wrap_text(line, font_name, 10, width)
            for part in wrapped:
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 24 * mm, reset_page)
                pdf.drawString(x, y, part)
                y -= 5 * mm
        return y

    @staticmethod
    def _draw_executive_title(pdf: canvas.Canvas, font_name: str, title: str, y: float, x: float) -> float:
        pdf.setStrokeColor(HexColor("#b28b67"))
        pdf.setLineWidth(1)
        pdf.line(x, y + 1.5 * mm, x + 18 * mm, y + 1.5 * mm)
        pdf.setFillColor(HexColor("#6b4d2f"))
        pdf.setFont(font_name, 11)
        pdf.drawString(x + 20 * mm, y, title)
        return y - 6.5 * mm

    @classmethod
    def _draw_executive_paragraph(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        text: str,
        x: float,
        y: float,
        width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#3a3028"))
        pdf.setFont(font_name, 10)
        for line in cls._wrap_text(text, font_name, 10, width):
            pdf.drawString(x, y, line)
            y -= 5.1 * mm
        return y

    @classmethod
    def _draw_executive_list(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        lines: list[str],
        x: float,
        y: float,
        width: float,
        bullet: bool = True,
    ) -> float:
        pdf.setFillColor(HexColor("#4a3b2b"))
        pdf.setFont(font_name, 9)
        indent = 4 * mm if bullet else 0
        for raw_line in lines:
            wrapped = cls._wrap_text(str(raw_line), font_name, 9, width - indent)
            for index, part in enumerate(wrapped):
                prefix = "• " if bullet and index == 0 else ""
                pdf.drawString(x + (indent if bullet else 0), y, prefix + part)
                y -= 4.6 * mm
            y -= 1.2 * mm
        return y

    @classmethod
    def _draw_executive_block(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        pdf.setFillColor(HexColor("#3a3028"))
        pdf.setFont(font_name, 9)
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9, width - 6 * mm)
                for index, part in enumerate(wrapped):
                    if reset_page is not None:
                        y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                    prefix = "• " if index == 0 else "  "
                    pdf.drawString(x + 3 * mm, y, prefix + part)
                    y -= 4.6 * mm
                continue
            pdf.setFillColor(HexColor("#2f261d"))
            pdf.setFont(font_name, 10)
            for part in cls._wrap_text(line, font_name, 10, width):
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                pdf.drawString(x, y, part)
                y -= 5 * mm
            pdf.setFillColor(HexColor("#3a3028"))
            pdf.setFont(font_name, 9)
        return y

    @classmethod
    def _draw_edinburgh_side_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        lines: list[str],
        y: float,
        sidebar_width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#2e2f45"))
        pdf.setFont(font_name, 13)
        pdf.drawString(8 * mm, y, title)
        y -= 6 * mm
        pdf.setFillColor(HexColor("#4b4b58"))
        pdf.setFont(font_name, 9)
        for line in lines:
            for part in cls._wrap_text(str(line), font_name, 9, sidebar_width - 20 * mm):
                pdf.drawString(10 * mm, y, part)
                y -= 4.8 * mm
            y -= 1.4 * mm
        return y

    @classmethod
    def _draw_edinburgh_rating_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        items: list[str],
        y: float,
        sidebar_width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#2e2f45"))
        pdf.setFont(font_name, 13)
        pdf.drawString(8 * mm, y, title)
        y -= 7 * mm
        label_width = sidebar_width - 38 * mm
        for index, item in enumerate(items[:5]):
            pdf.setFillColor(HexColor("#4b4b58"))
            pdf.setFont(font_name, 9)
            label = str(item)
            label = cls._wrap_text(label, font_name, 9, label_width)[0]
            pdf.drawString(8 * mm, y, label)
            level = max(2, 5 - (index % 2))
            dot_x = sidebar_width - 22 * mm
            for dot in range(5):
                pdf.setFillColor(HexColor("#5f5e86") if dot < level else HexColor("#d7d8e2"))
                pdf.circle(dot_x + (dot * 5 * mm), y + 1 * mm, 1.5 * mm, fill=1, stroke=0)
            y -= 6 * mm
        return y

    @classmethod
    def _draw_edinburgh_intro(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        text: str,
        x: float,
        y: float,
        width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#2f313f"))
        pdf.setFont(font_name, 12)
        pdf.drawString(x, y, "PROFILE SUMMARY")
        y -= 7 * mm
        pdf.setFillColor(HexColor("#4b4b58"))
        pdf.setFont(font_name, 10)
        for line in cls._wrap_text(text, font_name, 10, width):
            pdf.drawString(x, y, line)
            y -= 5.1 * mm
        pdf.setStrokeColor(HexColor("#d6d7df"))
        pdf.setLineWidth(0.8)
        pdf.line(x, y + 1 * mm, x + width, y + 1 * mm)
        return y - 7 * mm

    @classmethod
    def _draw_royal_intro(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        text: str,
        x: float,
        y: float,
        width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#4d627b"))
        pdf.setFont(font_name, 10)
        for line in cls._wrap_text(text, font_name, 10, width):
            pdf.drawString(x, y, line)
            y -= 5.1 * mm
        pdf.setStrokeColor(HexColor("#d7dde7"))
        pdf.setLineWidth(0.8)
        pdf.line(x, y + 1 * mm, x + width, y + 1 * mm)
        return y - 7 * mm

    @classmethod
    def _draw_royal_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        if not text:
            return y
        if reset_page is not None:
            y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
        pdf.setFillColor(HexColor("#26384c"))
        pdf.setFont(font_name, 15)
        pdf.drawString(x, y, title)
        y -= 6 * mm
        pdf.setStrokeColor(HexColor("#d7dde7"))
        pdf.line(x, y + 2 * mm, x + width, y + 2 * mm)
        y -= 2 * mm
        timeline_x = x + 16 * mm
        text_x = x + 26 * mm
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if "|" in line and not line.startswith("- "):
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                parts = [part.strip() for part in line.split("|") if part.strip()]
                left = parts[-1] if len(parts) > 1 else ""
                right = " | ".join(parts[:-1]) if len(parts) > 1 else parts[0]
                pdf.setStrokeColor(HexColor("#7e8da3"))
                pdf.line(timeline_x, y - 7 * mm, timeline_x, y + 2 * mm)
                pdf.setFillColor(HexColor("#ffffff"))
                pdf.circle(timeline_x, y, 1.7 * mm, fill=1, stroke=1)
                pdf.setFillColor(HexColor("#6d7f98"))
                pdf.setFont(font_name, 8.5)
                if left:
                    pdf.drawString(x, y, left)
                pdf.setFillColor(HexColor("#2e3d4d"))
                pdf.setFont(font_name, 10.5)
                for idx, part in enumerate(cls._wrap_text(right, font_name, 10.5, width - 30 * mm)):
                    pdf.drawString(text_x, y - idx * 4.8 * mm, part)
                y -= max(7 * mm, (len(cls._wrap_text(right, font_name, 10.5, width - 30 * mm)) * 4.8 * mm) + 1.5 * mm)
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9, width - 30 * mm)
                pdf.setFillColor(HexColor("#4c5f77"))
                pdf.setFont(font_name, 9)
                for idx, part in enumerate(wrapped):
                    if reset_page is not None:
                        y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                    prefix = "• " if idx == 0 else "  "
                    pdf.drawString(text_x + 3 * mm, y, prefix + part)
                    y -= 4.6 * mm
                continue
            pdf.setFillColor(HexColor("#4c5f77"))
            pdf.setFont(font_name, 9.2)
            for part in cls._wrap_text(line, font_name, 9.2, width - 30 * mm):
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                pdf.drawString(text_x, y, part)
                y -= 4.6 * mm
        return y - 5 * mm

    @classmethod
    def _draw_burgundy_side_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        lines: list[str],
        y: float,
        sidebar_width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont(font_name, 12)
        pdf.drawString(8 * mm, y, title)
        y -= 6 * mm
        pdf.setFont(font_name, 8.7)
        for line in lines:
            for part in cls._wrap_text(str(line), font_name, 8.7, sidebar_width - 16 * mm):
                pdf.drawString(8 * mm, y, part)
                y -= 4.8 * mm
            y -= 1.2 * mm
        return y

    @classmethod
    def _draw_burgundy_bars(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        items: list[str],
        y: float,
        sidebar_width: float,
        bar_color: str = "#ffffff",
        text_color: str = "#ffffff",
        title_color: str = "#ffffff",
    ) -> float:
        pdf.setFillColor(HexColor(title_color))
        pdf.setFont(font_name, 12)
        pdf.drawString(8 * mm, y, title)
        y -= 7 * mm
        pdf.setFont(font_name, 8.7)
        for idx, item in enumerate(items):
            pdf.setFillColor(HexColor(text_color))
            pdf.drawString(8 * mm, y, str(item))
            pdf.setFillColor(HexColor(bar_color))
            pdf.rect(8 * mm, y - 2.6 * mm, sidebar_width - 24 * mm - (idx % 2) * 6 * mm, 1.3 * mm, fill=1, stroke=0)
            y -= 7 * mm
        return y

    @classmethod
    def _draw_burgundy_card(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        if reset_page is not None:
            y = cls._ensure_canvas_space(pdf, y, 32 * mm, reset_page)
        pdf.setFillColor(HexColor("#6b2c3e"))
        pdf.rect(x, y, width, 8 * mm, fill=1, stroke=0)
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont(font_name, 11)
        pdf.drawString(x + 5 * mm, y + 2.2 * mm, title)
        card_y = y - 3 * mm
        pdf.setFillColor(HexColor("#ffffff"))
        content_height = max(26 * mm, min(58 * mm, (len(cls._normalize_content(text)) + 2) * 6 * mm))
        pdf.rect(x, card_y - content_height, width, content_height, fill=1, stroke=1)
        inner_y = card_y - 6 * mm
        pdf.setFillColor(HexColor("#2f2f35"))
        pdf.setFont(font_name, 9)
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                inner_y -= 2 * mm
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9, width - 16 * mm)
                for idx, part in enumerate(wrapped):
                    pdf.drawString(x + 8 * mm, inner_y, ("• " if idx == 0 else "  ") + part)
                    inner_y -= 4.6 * mm
                continue
            for part in cls._wrap_text(line, font_name, 9, width - 12 * mm):
                pdf.drawString(x + 5 * mm, inner_y, part)
                inner_y -= 4.8 * mm
        return card_y - content_height - 6 * mm

    @classmethod
    def _draw_onyx_side_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        lines: list[str],
        y: float,
        sidebar_width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont(font_name, 12)
        pdf.drawString(6 * mm, y, title)
        y -= 6 * mm
        pdf.setFont(font_name, 8.7)
        for line in lines:
            for part in cls._wrap_text(str(line), font_name, 8.7, sidebar_width - 14 * mm):
                pdf.drawString(6 * mm, y, part)
                y -= 4.8 * mm
            y -= 1 * mm
        return y

    @classmethod
    def _draw_onyx_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        if reset_page is not None:
            y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
        pdf.setFillColor(HexColor("#222222"))
        pdf.setFont(font_name, 13)
        pdf.drawString(x, y, title)
        y -= 3 * mm
        pdf.setStrokeColor(HexColor("#222222"))
        pdf.line(x, y, x + width, y)
        y -= 5 * mm
        pdf.setFillColor(HexColor("#2d2d2d"))
        pdf.setFont(font_name, 9.3)
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9.3, width - 8 * mm)
                for idx, part in enumerate(wrapped):
                    if reset_page is not None:
                        y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                    pdf.drawString(x + 3 * mm, y, ("• " if idx == 0 else "  ") + part)
                    y -= 4.8 * mm
                continue
            for part in cls._wrap_text(line, font_name, 9.3, width):
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                pdf.drawString(x, y, part)
                y -= 4.8 * mm
        return y - 5 * mm

    @classmethod
    def _draw_classic_header_meta(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        profile: dict,
        x: float,
        y: float,
        width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#444444"))
        pdf.setFont(font_name, 9.5)
        details = [
            ("Phone", profile.get("phone")),
            ("Email", profile.get("email")),
            ("Location", profile.get("location")),
        ]
        details = [(label, value) for label, value in details if value]
        row_y = y
        for label, value in details:
            pdf.drawString(x, row_y, label)
            pdf.drawString(x + 38 * mm, row_y, str(value))
            row_y -= 6 * mm
        pdf.setStrokeColor(HexColor("#d1d1d1"))
        pdf.line(x, row_y, x + width, row_y)
        return row_y - 8 * mm

    @classmethod
    def _draw_classic_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        if reset_page is not None:
            y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
        pdf.setFillColor(HexColor("#2b2b2b"))
        pdf.setFont(font_name, 14)
        pdf.drawString(x, y, title)
        y -= 4 * mm
        pdf.setStrokeColor(HexColor("#d1d1d1"))
        pdf.line(x, y, x + width, y)
        y -= 6 * mm
        pdf.setFillColor(HexColor("#4b4b4b"))
        pdf.setFont(font_name, 9.2)
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9.2, width - 8 * mm)
                for idx, part in enumerate(wrapped):
                    if reset_page is not None:
                        y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                    pdf.drawString(x + 3 * mm, y, ("• " if idx == 0 else "  ") + part)
                    y -= 4.8 * mm
                continue
            for part in cls._wrap_text(line, font_name, 9.2, width):
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                pdf.drawString(x, y, part)
                y -= 4.8 * mm
        return y - 5 * mm

    @classmethod
    def _draw_classic_timeline_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page,
    ) -> float:
        y = cls._draw_classic_section(pdf, font_name, title, "", x, y, width, reset_page=reset_page)
        timeline_x = x + 34 * mm
        text_x = x + 44 * mm
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if "|" in line and not line.startswith("- "):
                y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                parts = [part.strip() for part in line.split("|") if part.strip()]
                date_text = parts[-1] if len(parts) > 1 else ""
                main_text = " | ".join(parts[:-1]) if len(parts) > 1 else parts[0]
                pdf.setStrokeColor(HexColor("#3d3d3d"))
                pdf.line(timeline_x, y - 7 * mm, timeline_x, y + 1 * mm)
                pdf.circle(timeline_x, y, 1.3 * mm, fill=1, stroke=0)
                pdf.setFillColor(HexColor("#555555"))
                pdf.setFont(font_name, 8.7)
                if date_text:
                    pdf.drawString(x, y, date_text)
                pdf.setFillColor(HexColor("#2b2b2b"))
                pdf.setFont(font_name, 10)
                for idx, part in enumerate(cls._wrap_text(main_text, font_name, 10, width - 48 * mm)):
                    pdf.drawString(text_x, y - idx * 4.8 * mm, part)
                y -= max(7 * mm, len(cls._wrap_text(main_text, font_name, 10, width - 48 * mm)) * 4.8 * mm + 1 * mm)
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9, width - 48 * mm)
                pdf.setFillColor(HexColor("#4b4b4b"))
                pdf.setFont(font_name, 9)
                for idx, part in enumerate(wrapped):
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                    pdf.drawString(text_x + 3 * mm, y, ("• " if idx == 0 else "  ") + part)
                    y -= 4.8 * mm
                continue
        return y - 6 * mm

    @classmethod
    def _draw_crimson_side_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        lines: list[str],
        x: float,
        y: float,
        width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont(font_name, 11)
        pdf.drawString(x, y, title)
        y -= 7 * mm
        pdf.setFont(font_name, 8.6)
        for line in lines:
            for part in cls._wrap_text(str(line), font_name, 8.6, width):
                pdf.drawString(x, y, part)
                y -= 4.8 * mm
            y -= 1.2 * mm
        return y

    @classmethod
    def _draw_crimson_skill_dots(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        items: list[str],
        x: float,
        y: float,
        width: float,
    ) -> float:
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont(font_name, 11)
        pdf.drawString(x, y, title)
        y -= 7 * mm
        for idx, item in enumerate(items):
            pdf.setFont(font_name, 8.6)
            pdf.drawString(x, y, item)
            level = max(2, 5 - (idx % 2))
            dot_x = x
            y -= 4.5 * mm
            for dot in range(5):
                pdf.setFillColor(HexColor("#ffffff") if dot < level else HexColor("#f1a0a0"))
                pdf.circle(dot_x + (dot * 4.2 * mm), y, 1.1 * mm, fill=1, stroke=0)
            y -= 5 * mm
        return y

    @classmethod
    def _draw_crimson_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page,
    ) -> float:
        y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
        pdf.setFillColor(HexColor("#2d2d2d"))
        pdf.setFont(font_name, 14)
        pdf.drawString(x, y, title)
        y -= 4 * mm
        pdf.setStrokeColor(HexColor("#c9c9c9"))
        pdf.line(x, y, x + width, y)
        y -= 6 * mm
        pdf.setFillColor(HexColor("#333333"))
        pdf.setFont(font_name, 9.3)
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9.3, width - 8 * mm)
                for idx, part in enumerate(wrapped):
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                    pdf.drawString(x + 3 * mm, y, ("• " if idx == 0 else "  ") + part)
                    y -= 4.8 * mm
                continue
            for part in cls._wrap_text(line, font_name, 9.3, width):
                y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                pdf.drawString(x, y, part)
                y -= 4.8 * mm
        return y - 5 * mm

    @classmethod
    def _draw_edinburgh_section(
        cls,
        pdf: canvas.Canvas,
        font_name: str,
        title: str,
        text: str,
        x: float,
        y: float,
        width: float,
        reset_page=None,
    ) -> float:
        if not text:
            return y
        pdf.setFillColor(HexColor("#2f313f"))
        pdf.setFont(font_name, 16)
        if reset_page is not None:
            y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
        pdf.drawString(x, y, title)
        y -= 7 * mm
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 2 * mm
                continue
            if line.startswith("- "):
                wrapped = cls._wrap_text(line[2:], font_name, 9, width - 7 * mm)
                pdf.setFillColor(HexColor("#4b4b58"))
                pdf.setFont(font_name, 9)
                for idx, part in enumerate(wrapped):
                    if reset_page is not None:
                        y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                    prefix = "• " if idx == 0 else "  "
                    pdf.drawString(x + 3 * mm, y, prefix + part)
                    y -= 4.8 * mm
                continue
            if "|" in line:
                left, right = [part.strip() for part in line.split("|", 1)]
                pdf.setFillColor(HexColor("#2f313f"))
                pdf.setFont(font_name, 11)
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                pdf.drawString(x, y, left)
                text_width = pdfmetrics.stringWidth(right, font_name, 9)
                pdf.setFillColor(HexColor("#6a6e7c"))
                pdf.setFont(font_name, 9)
                pdf.drawString(x + width - text_width, y, right)
                y -= 5.2 * mm
                continue
            pdf.setFillColor(HexColor("#4b4b58"))
            pdf.setFont(font_name, 9)
            for part in cls._wrap_text(line, font_name, 9, width):
                if reset_page is not None:
                    y = cls._ensure_canvas_space(pdf, y, 26 * mm, reset_page)
                pdf.drawString(x, y, part)
                y -= 4.9 * mm
        pdf.setStrokeColor(HexColor("#d6d7df"))
        pdf.setLineWidth(0.8)
        pdf.line(x, y + 1 * mm, x + width, y + 1 * mm)
        return y - 7 * mm

    @staticmethod
    def _resolve_photo_path(photo_value: str | None) -> str | None:
        if not photo_value:
            return None
        candidate = Path(photo_value)
        if candidate.is_file():
            studio_candidate = ResumeService._resolve_studio_variant(candidate)
            return studio_candidate or str(candidate)
        raw = photo_value.replace("\\", "/").lstrip("/")
        base = Path(__file__).resolve().parents[3]
        for possible in (
            base / raw,
            base / "frontend" / "public" / "photos" / Path(raw).name,
            base / "assets" / "profile_photos" / Path(raw).name,
        ):
            if possible.is_file():
                studio_candidate = ResumeService._resolve_studio_variant(possible)
                return studio_candidate or str(possible)
        return None

    @staticmethod
    def _resolve_studio_variant(photo_path: Path) -> str | None:
        base = Path(__file__).resolve().parents[3]
        processed = base / "frontend" / "public" / "photos" / "processed" / f"{photo_path.stem}-studio.png"
        if processed.is_file():
            return str(processed)
        return None

    @classmethod
    def _draw_profile_photo(
        cls,
        pdf: canvas.Canvas,
        photo_path: str | None,
        center_x: float,
        center_y: float,
        radius: float,
        placeholder_color: str,
    ) -> float:
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.circle(center_x, center_y, radius, fill=1, stroke=0)
        if photo_path:
            try:
                image = ImageReader(photo_path)
                size = radius * 2
                pdf.saveState()
                clip = pdf.beginPath()
                clip.circle(center_x, center_y, radius - 1.2 * mm)
                pdf.clipPath(clip, stroke=0, fill=0)
                pdf.drawImage(
                    image,
                    center_x - (size / 2),
                    center_y - (size / 2),
                    width=size,
                    height=size,
                    preserveAspectRatio=False,
                    mask="auto",
                )
                pdf.restoreState()
                pdf.setStrokeColor(HexColor("#ffffff"))
                pdf.setLineWidth(1.6)
                pdf.circle(center_x, center_y, radius - 0.4 * mm, fill=0, stroke=1)
                return center_y - radius
            except Exception:
                pass
        pdf.setFillColor(HexColor(placeholder_color))
        pdf.circle(center_x, center_y, radius - 2 * mm, fill=1, stroke=0)
        pdf.setFillColor(HexColor("#7d8097"))
        pdf.circle(center_x, center_y + 3 * mm, 6 * mm, fill=1, stroke=0)
        pdf.roundRect(center_x - 10 * mm, center_y - 14 * mm, 20 * mm, 12 * mm, 5 * mm, fill=1, stroke=0)
        return center_y - radius

    @staticmethod
    def _ensure_canvas_space(pdf: canvas.Canvas, y: float, min_y: float, reset_page) -> float:
        if y >= min_y:
            return y
        pdf.showPage()
        return reset_page()

    @staticmethod
    def _split_sections(content: str, profile: dict | None = None) -> dict[str, str]:
        sections: dict[str, list[str]] = {}
        current_heading = "EXPERIENCE"
        sections[current_heading] = []
        for raw_line in content.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if line.isupper() and len(line) < 40:
                current_heading = line
                sections.setdefault(current_heading, [])
                continue
            sections.setdefault(current_heading, []).append(line)
        combined = {key: "\n".join(values).strip() for key, values in sections.items() if values}
        if not profile:
            return combined

        identity_lines = {
            str(profile.get("full_name") or "").strip(),
            str(profile.get("target_role") or "").strip(),
            str(profile.get("email") or "").strip(),
            str(profile.get("phone") or "").strip(),
            str(profile.get("location") or "").strip(),
        }
        identity_lines = {line for line in identity_lines if line and line.lower() != "none"}

        cleaned: dict[str, str] = {}
        for heading, block in combined.items():
            filtered_lines = [
                line for line in block.splitlines()
                if line.strip() and line.strip() not in identity_lines
            ]
            if filtered_lines:
                cleaned[heading] = "\n".join(filtered_lines).strip()
        return cleaned

    @staticmethod
    def _collect_links(profile: dict) -> list[str]:
        links = []
        for value in (
            profile.get("linkedin_url"),
            profile.get("github_url"),
            profile.get("portfolio_url"),
        ):
            if value:
                links.append(value)
        return links

    @staticmethod
    def _group_skills(skills: list[str]) -> dict[str, list[str]]:
        if not skills:
            return {}

        buckets: dict[str, list[str]] = {
            "Frontend": [],
            "Backend": [],
            "Database": [],
            "DevOps": [],
            "Tools": [],
        }
        for skill in skills:
            lowered = skill.lower()
            if lowered in {"react", "next.js", "typescript", "javascript", "tailwind css", "redux", "html", "css"}:
                buckets["Frontend"].append(skill)
            elif lowered in {"python", "fastapi", "django", "node.js", "express", "graphql", "rest api", "rest"}:
                buckets["Backend"].append(skill)
            elif lowered in {"postgresql", "mysql", "mongodb", "redis", "sqlite"}:
                buckets["Database"].append(skill)
            elif lowered in {"docker", "kubernetes", "aws", "ci/cd", "nginx"}:
                buckets["DevOps"].append(skill)
            else:
                buckets["Tools"].append(skill)
        return {key: value for key, value in buckets.items() if value}

    @staticmethod
    def _escape_pdf(text: str) -> str:
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )

    @staticmethod
    def _slugify(value: str) -> str:
        clean = "".join(ch.lower() if ch.isalnum() else "-" for ch in value)
        while "--" in clean:
            clean = clean.replace("--", "-")
        return clean.strip("-") or "cv"

    @staticmethod
    def _ensure_pdf_font() -> str:
        font_name = "ArialUnicodeSafe"
        if font_name in pdfmetrics.getRegisteredFontNames():
            return font_name

        for candidate in (
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/calibri.ttf"),
            Path("C:/Windows/Fonts/segoeui.ttf"),
        ):
            if candidate.exists():
                pdfmetrics.registerFont(TTFont(font_name, str(candidate)))
                return font_name
        return "Helvetica"
